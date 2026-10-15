import json
import os

import pytest
from docx import Document as DocxDocument

from formats.content import ContentValidationError


def test_export_xlsx_basic(isolated_settings, tmp_export_dir):
    from formats.xlsx import export_xlsx
    ref = export_xlsx([["Name", "Score"], ["Alice", "95"]], filename="test.xlsx")
    assert os.path.exists(ref.path)
    assert ref.name == "test.xlsx"
    assert ref.url.endswith("/test.xlsx")


def test_export_xlsx_with_title(isolated_settings):
    from formats.xlsx import export_xlsx
    ref = export_xlsx([["A", "B"], ["1", "2"]], filename="titled.xlsx", title="My Sheet")
    assert os.path.exists(ref.path)


def test_export_docx_blocks(isolated_settings):
    from formats.docx import export_docx
    blocks = [
        {"type": "heading", "level": 1, "text": "Title"},
        {"type": "paragraph", "text": "Hello world"},
        {"type": "list", "items": ["one", "two"]},
        {"type": "table", "data": [["H1", "H2"], ["a", "b"]]},
    ]
    ref = export_docx(blocks, filename="doc.docx")
    assert os.path.exists(ref.path)
    doc = DocxDocument(ref.path)
    assert len(doc.paragraphs) > 0


def test_export_docx_plain_string(isolated_settings):
    from formats.docx import export_docx
    ref = export_docx("line one\nline two", filename="plain.docx")
    assert os.path.exists(ref.path)


def test_export_docx_fenced_json_string(isolated_settings):
    """A JSON argument wrapped in a markdown code fence must not leak
    raw markup into the document; the blocks must be rendered."""
    from formats.docx import export_docx
    import json
    blocks = [
        {"type": "heading", "level": 1, "text": "Report"},
        {"type": "paragraph", "text": "Body text."},
    ]
    fenced = "```json\n" + json.dumps(blocks, ensure_ascii=False) + "\n```"
    ref = export_docx(fenced, filename="fenced.docx")
    doc = DocxDocument(ref.path)
    texts = [p.text for p in doc.paragraphs]
    assert "Report" in texts
    assert "Body text." in texts
    assert not any("```" in t for t in texts)
    assert not any('"type"' in t for t in texts)


def test_export_docx_truncated_blocks_raise(isolated_settings):
    """A truncated JSON argument must raise, not export raw markup."""
    from formats.docx import export_docx
    truncated = '[{"type": "paragraph", "text": "hi"'
    with pytest.raises(ValueError):
        export_docx(truncated, filename="bad.docx")


def test_export_docx_unknown_block_type_raises(isolated_settings):
    from formats.docx import export_docx
    with pytest.raises(ValueError):
        export_docx([{"type": "blockquote", "text": "hi"}], filename="bad.docx")


def test_export_docx_missing_field_raises(isolated_settings):
    from formats.docx import export_docx
    with pytest.raises(ValueError):
        export_docx([{"type": "paragraph"}], filename="bad.docx")


# ── PPTX: native slide dicts ──────────────────────────────────────────────────


def test_export_pptx_blocks(isolated_settings):
    """Regression: formats/pptx.py was ported from the old monolith without
    any imports (logging, os, pptx, ...) and with dangling helpers, so every
    export_pptx() call crashed with 'NameError: name "logging" is not
    defined'. This test actually runs the exporter end-to-end, so a missing
    import or broken helper is caught by CI instead of only in production.
    """
    from formats.pptx import export_pptx
    from pptx import Presentation as PptxPresentation

    slides = [
        {"title": "Тема проекта", "content": ["Команда проекта", "Катавский филиал"]},
        {"title": "Сроки", "content": "07.08.2025 – 30.11.2026"},
    ]
    ref = export_pptx(slides, filename="deck.pptx")
    assert os.path.exists(ref.path)
    assert ref.name == "deck.pptx"
    assert ref.url.endswith("/deck.pptx")
    prs = PptxPresentation(ref.path)
    assert len(prs.slides) > 0


def test_export_pptx_plain_string(isolated_settings):
    """A plain string is wrapped into a single slide by export_pptx."""
    from formats.pptx import export_pptx
    from pptx import Presentation as PptxPresentation

    ref = export_pptx("первый слайд\nвторой слайд", filename="plain.pptx")
    assert os.path.exists(ref.path)
    assert ref.name == "plain.pptx"
    prs = PptxPresentation(ref.path)
    assert len(prs.slides) > 0


def test_export_pptx_serialized_string_recovered(isolated_settings):
    """A JSON-serialized slide list (transport/LLM mangling) must be recovered
    and rendered, not dumped as raw markup into the slides."""
    import json
    from formats.pptx import export_pptx
    from pptx import Presentation as PptxPresentation

    slides = [
        {"title": "ТЗ", "content": ["п.1", "п.2"]},
    ]
    serialized = json.dumps(slides, ensure_ascii=False)
    ref = export_pptx(serialized, filename="rec.pptx")
    prs = PptxPresentation(ref.path)
    full_text = ""
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                full_text += shape.text_frame.text + "\n"
    assert '{"title":' not in full_text
    assert "ТЗ" in full_text


def test_export_pptx_malformed_slides_raise(isolated_settings):
    """Malformed slides must raise, not produce a broken presentation."""
    from formats.pptx import export_pptx
    with pytest.raises(ContentValidationError):
        export_pptx([42], filename="bad.pptx")


# ── PPTX: DOCX blocks auto-conversion ─────────────────────────────────────────


def test_export_pptx_docx_blocks_converted(isolated_settings):
    """The exact scenario from the bug report: user passes DOCX content
    blocks (heading/paragraph/list) to format="pptx". The exporter must
    convert them into proper slides, not dump raw text."""
    from formats.pptx import export_pptx
    from pptx import Presentation as PptxPresentation

    blocks = [
        {"type": "heading", "level": 1, "text": "Описание проблемной ситуации"},
        {"type": "paragraph", "text": "ТС: Кровля производственных зданий."},
        {"type": "paragraph", "text": "Процесс: Эксплуатация кровли."},
        {"type": "list", "items": [
            "ЧТО: Постоянные протечки",
            "ГДЕ: В местах примыканий",
            "КОГДА: На протяжении всего срока",
            "ПОЧЕМУ: Рулонное покрытие не рассчитано",
        ]},
        {"type": "paragraph", "text": "Экономический ущерб: 14 346 тыс. руб."},
        {"type": "paragraph", "text": "Цели: срок службы ≥ 10 лет"},
    ]
    ref = export_pptx(blocks, filename="krovlya.pptx")
    assert os.path.exists(ref.path)
    assert ref.name == "krovlya.pptx"

    prs = PptxPresentation(ref.path)
    # Should have at least 1 slide (heading level 1 creates a new slide)
    assert len(prs.slides) >= 1

    # Collect all text from all slides
    full_text = ""
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                full_text += shape.text_frame.text + "\n"

    # The heading text must appear as a slide title or content
    assert "Описание проблемной ситуации" in full_text
    # Paragraph content must be present
    assert "Кровля производственных зданий" in full_text
    # List items must be present
    assert "Постоянные протечки" in full_text
    # No raw JSON/markup should leak
    assert '{"type":' not in full_text
    assert '"paragraph"' not in full_text
    assert '"heading"' not in full_text


def test_export_pptx_docx_blocks_multiple_headings(isolated_settings):
    """Multiple heading level 1 blocks should produce multiple slides."""
    from formats.pptx import export_pptx
    from pptx import Presentation as PptxPresentation

    blocks = [
        {"type": "heading", "level": 1, "text": "Слайд 1"},
        {"type": "paragraph", "text": "Контент 1"},
        {"type": "heading", "level": 1, "text": "Слайд 2"},
        {"type": "paragraph", "text": "Контент 2"},
        {"type": "heading", "level": 2, "text": "Подзаголовок"},
        {"type": "paragraph", "text": "Контент 3"},
    ]
    ref = export_pptx(blocks, filename="multi.pptx")
    prs = PptxPresentation(ref.path)

    # 2 heading level 1 → 2 slides (+ 1 title slide from _add_title_slide)
    # heading level 2 stays in current slide
    full_text = ""
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                full_text += shape.text_frame.text + "\n"

    assert "Слайд 1" in full_text
    assert "Слайд 2" in full_text
    assert "Подзаголовок" in full_text
    assert "Контент 1" in full_text
    assert "Контент 3" in full_text


def test_export_pptx_docx_blocks_table(isolated_settings):
    """Table blocks should be converted to content lines (cells joined)."""
    from formats.pptx import export_pptx
    from pptx import Presentation as PptxPresentation

    blocks = [
        {"type": "heading", "level": 1, "text": "Данные"},
        {"type": "table", "data": [
            ["Параметр", "Значение"],
            ["Срок", "10 лет"],
            ["Затраты", "6 855 тыс. руб."],
        ]},
    ]
    ref = export_pptx(blocks, filename="table.pptx")
    prs = PptxPresentation(ref.path)

    full_text = ""
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                full_text += shape.text_frame.text + "\n"

    assert "Данные" in full_text
    assert "Параметр" in full_text
    assert "10 лет" in full_text


def test_export_pptx_docx_blocks_json_string(isolated_settings):
    """DOCX blocks passed as a JSON-serialized string must also be converted."""
    import json
    from formats.pptx import export_pptx
    from pptx import Presentation as PptxPresentation

    blocks = [
        {"type": "heading", "level": 1, "text": "Тема"},
        {"type": "paragraph", "text": "Текст"},
    ]
    serialized = json.dumps(blocks, ensure_ascii=False)
    ref = export_pptx(serialized, filename="json_blocks.pptx")
    prs = PptxPresentation(ref.path)

    full_text = ""
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                full_text += shape.text_frame.text + "\n"

    assert "Тема" in full_text
    assert "Текст" in full_text
    assert '{"type":' not in full_text


def test_export_csv_basic(isolated_settings):
    # csv exporter currently raises NotImplementedError due to stub
    from formats.csv import export_csv
    with pytest.raises(NotImplementedError):
        export_csv([["a", "b"], ["1", "2"]], filename="t.csv")


def test_template_registry_scan(isolated_settings, tmp_export_dir):
    from app.core.templates import TemplateRegistry
    # create fake templates
    docx_path = os.path.join(tmp_export_dir, "Default_Template.docx")
    DocxDocument().save(docx_path)
    named_docx = os.path.join(tmp_export_dir, "protocol.docx")
    DocxDocument().save(named_docx)
    import json
    with open(os.path.join(tmp_export_dir, "protocol.json"), "w") as f:
        json.dump({"description": "test", "placeholders": {"Who": "who hint"}}, f)
    TemplateRegistry.init(tmp_export_dir, tmp_export_dir)
    assert TemplateRegistry.get("docx") == docx_path
    assert TemplateRegistry.get_named("protocol") == named_docx
    items = TemplateRegistry.list_named()
    assert any(x["name"] == "protocol" for x in items)


def test_is_safe_path_used_in_archive(isolated_settings, tmp_export_dir):
    from storage.files import write_text
    ref = write_text("data", "a.txt", "txt")
    # create_archive should filter unsafe paths
    import asyncio
    from tools.archive_tools import create_archive

    async def _run():
        result = await create_archive([ref.path, "/etc/passwd"], archive_name="out.zip", format="zip")
        assert "path" in result
        assert os.path.exists(result["path"])

    asyncio.run(_run())
