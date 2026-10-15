"""Tests for structured-input validation: raw markup must not reach the file.

These tests cover the user-visible bug: the caller sent content blocks,
the transport serialized them to a JSON string (with prose, code fences or a
byte order mark around them), and the old code wrote the raw markup into the
.docx as if it were plain text.
"""
from __future__ import annotations

import json

from formats.coerce import coerce
from formats.validate import normalize_structured_input

BLOCKS = [
    {"type": "heading", "level": 1, "text": "Test"},
    {"type": "paragraph", "text": "Body"},
    {"type": "table", "data": [["A", "B"], ["1", "2"]]},
]


# ── coerce: recovery of JSON strings ───────────────────────────────────────

def test_coerce_clean_json_string():
    assert coerce(json.dumps(BLOCKS)) == BLOCKS


def test_coerce_json_string_with_bom():
    assert coerce("\ufeff" + json.dumps(BLOCKS)) == BLOCKS


def test_coerce_json_string_in_code_fence():
    fenced = "```json\n" + json.dumps(BLOCKS) + "\n```"
    assert coerce(fenced) == BLOCKS


def test_coerce_json_string_with_prose_prefix():
    prefixed = "Here is the content: " + json.dumps(BLOCKS)
    assert coerce(prefixed) == BLOCKS


def test_coerce_json_string_with_trailing_prose():
    s = json.dumps(BLOCKS) + " Let me know if you need changes."
    assert coerce(s) == BLOCKS


def test_coerce_plain_text_with_brackets_not_json_stays_unchanged():
    s = "Use [ brackets ] in text, not JSON"
    assert coerce(s) == s


# ── normalize_structured_input: valid inputs pass ──────────────────────────

def test_normalize_valid_block_list_passes():
    value, err = normalize_structured_input(BLOCKS)
    assert err is None
    assert value == BLOCKS


def test_normalize_plain_string_passes():
    value, err = normalize_structured_input("line one\nline two")
    assert err is None
    assert value == "line one\nline two"


def test_normalize_list_with_strings_and_blocks_passes():
    value, err = normalize_structured_input(["free text", {"type": "paragraph", "text": "ok"}])
    assert err is None


def test_normalize_json_string_recovers_blocks():
    value, err = normalize_structured_input(json.dumps(BLOCKS))
    assert err is None
    assert value == BLOCKS


def test_normalize_one_item_json_string_list_recovers_blocks():
    value, err = normalize_structured_input([json.dumps(BLOCKS)])
    assert err is None
    assert value == BLOCKS


def test_normalize_wrapper_dict_recovers_blocks():
    wrapped = {"title": "Report", "content": BLOCKS}
    value, err = normalize_structured_input(wrapped)
    assert err is None
    assert value == BLOCKS


def test_normalize_fenced_json_string_recovers_blocks():
    fenced = "```json\n" + json.dumps(BLOCKS) + "\n```"
    value, err = normalize_structured_input(fenced)
    assert err is None
    assert value == BLOCKS


# ── normalize_structured_input: invalid input fails with a clear error ─────

def test_normalize_unknown_block_type_fails():
    value, err = normalize_structured_input([{"type": "video", "url": "x"}])
    assert err is not None
    assert err.code == "INVALID_INPUT"
    assert "video" in err.message
    assert "paragraph" in err.message  # valid types listed in the error


def test_normalize_table_without_data_fails():
    value, err = normalize_structured_input([{"type": "table"}])
    assert err is not None
    assert err.code == "INVALID_INPUT"


def test_normalize_list_without_items_fails():
    value, err = normalize_structured_input([{"type": "list", "text": "oops"}])
    assert err is not None
    assert err.code == "INVALID_INPUT"


def test_normalize_non_dict_item_fails():
    value, err = normalize_structured_input([42])
    assert err is not None
    assert err.code == "INVALID_INPUT"


def test_normalize_garbage_dict_fails():
    value, stage_err = normalize_structured_input([{"foo": "bar"}])
    assert stage_err is not None
    assert stage_err.code == "INVALID_INPUT"


# ── end-to-end: the produced .docx must not contain raw markup ─────────────

def test_docx_from_json_string_has_no_raw_markup(isolated_settings):
    """The user-visible bug: blocks arrive as a JSON string, the docx must
    render them as headings/tables, not contain the raw markup."""
    from formats.docx import export_docx
    from docx import Document as DocxDocument

    value, err = normalize_structured_input(json.dumps(BLOCKS), "docx")
    assert err is None
    ref = export_docx(value, filename="no_leak.docx")

    doc = DocxDocument(ref.path)
    all_text = "\n".join(p.text for p in doc.paragraphs)
    assert '{"type"' not in all_text
    assert len(doc.tables) == 1
    assert doc.paragraphs[0].text == "Test"


def test_docx_from_prose_wrapped_json_has_no_raw_markup(isolated_settings):
    """A JSON string with prose around it (typical model output) must also
    be recovered, not written to the file as-is."""
    from formats.docx import export_docx
    from docx import Document as DocxDocument

    payload = "Вот содержимое документа: " + json.dumps(BLOCKS, ensure_ascii=False)
    value, err = normalize_structured_input(payload, "docx")
    assert err is None
    ref = export_docx(value, filename="prose_recovered.docx")

    doc = DocxDocument(ref.path)
    all_text = "\n".join(p.text for p in doc.paragraphs)
    assert '{"type"' not in all_text
    assert len(doc.tables) == 1
